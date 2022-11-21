# David Renan
# Pedro Corrêa

# Importação de módulos
import openpyxl
import time
from selenium import webdriver
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By

# Seleciona o driver mais atualizado do webdriver e instala para previnir erros de atualização e versão do navegador.
servico = Service(ChromeDriverManager().install())
navegador = webdriver.Chrome(service=servico)

# Acessando site
navegador.get('https://www.google.com/')

#Acessando a busca
navegador.find_element(By.NAME, 'q').send_keys('tempo agora Itatiaia-RJ')

# Tempo para carregamento da pagina completa
time.sleep(3)

# Acessando o botao
navegador.find_element(By.NAME, 'btnK').click()

# Encontrar temperatura
texto = navegador.find_element(By.ID, 'wob_tm').text

print ("A temperatura é: ", texto)

# ------------------------------------------------------------------------------------------------

# Criando planilha
Plan = openpyxl.Workbook()

# Criando pagina especifico
Plan.create_sheet('Temperaturas')

# Selecionando uma página
pag_tempo = Plan['Temperaturas']

# Inserindo registros
pag_tempo.append(['Data', 'Temperatura'])

# Salvar planilha
Plan.save ('Planilha de temperaturas.xlsx')
time.sleep(5)

# Carregando planilha
Carreg_Plan = openpyxl.load_workbook('Planilha de temperaturas.xlsx')
# selecionando aba da planilha
Selec_tempo = Plan['Temperaturas']